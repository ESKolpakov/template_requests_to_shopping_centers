import zipfile

import docx
from fixtures import build_synthetic_template

from uktc_letters.docx_template import RowSpec, render_letter


def _base_context():
    return {
        "RECIPIENT_BLOCK": "Администрация торгового центра\nул. Примерная, д.1",
        "GREETING": "",
        "TC_NAME": "Пример",
        "TC_ADDRESS": "ул. Примерная, д.1",
        "DEADLINE_DATE": "20.07.2026",
        "SUBJECT_WORD_FORM": "хозяйствующем субъекте",
        "VERB_OSUSHESTVLYAL": "осуществлял",
        "EXECUTOR_NAME": "И.И.Тестов",
        "EXECUTOR_PHONE": "8 (495) 000-00-00",
        "SIGNATORY_NAME": "Тестов Т.Т.",
    }


def test_render_letter_simple_placeholders(tmp_path):
    tpl = tmp_path / "tpl.docx"
    build_synthetic_template(str(tpl))

    row_specs = [
        RowSpec(
            num="1", name="Магазин", subject='ООО "Компания"', inn="7700000001",
            address="ул. Примерная, д.1", floor="1", room="-", date="01.07.2026",
            object_id="1000000001",
        )
    ]
    out = tmp_path / "out.docx"
    found = render_letter(str(tpl), _base_context(), row_specs, str(out))

    assert "RECIPIENT_BLOCK" in found
    assert "TC_NAME" in found

    d = docx.Document(str(out))
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "{{" not in body_text
    assert "торговый центр «Пример»" in body_text
    assert "осуществлял хозяйствующем субъекте" in body_text  # см. базовый контекст выше
    assert "20.07.2026" in body_text


def test_render_letter_multiline_recipient_block_creates_break(tmp_path):
    tpl = tmp_path / "tpl.docx"
    build_synthetic_template(str(tpl))
    row_specs = [
        RowSpec(
            num="1", name="Магазин", subject="X", inn="1", address="Y",
            floor="1", room="-", date="01.01.2026", object_id="1",
        )
    ]
    out = tmp_path / "out.docx"
    render_letter(str(tpl), _base_context(), row_specs, str(out))

    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "<w:br/>" in xml


def test_render_letter_rowspan_merging(tmp_path):
    tpl = tmp_path / "tpl.docx"
    build_synthetic_template(str(tpl))

    common = dict(
        name="Общий объект", address="ул. Общая, д.1", floor="1", room="-",
        date="01.01.2026", object_id="42",
    )
    row_specs = [
        RowSpec(num="1", subject="Иванов", inn="111", merge_start=True, **common),
        RowSpec(num="2", subject="Петров", inn="222", merge_continue=True, **common),
    ]
    out = tmp_path / "out.docx"
    render_letter(str(tpl), _base_context(), row_specs, str(out))

    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count('w:val="restart"') == 6  # 6 объединяемых столбцов
    assert xml.count("<w:vMerge/>") == 6

    d = docx.Document(str(out))
    obj_table = d.tables[1]
    assert len(obj_table.rows) == 3  # заголовок + 2 строки
    assert obj_table.rows[1].cells[2].text == "Иванов"
    assert obj_table.rows[2].cells[2].text == "Петров"


def test_render_letter_single_row_has_no_vmerge(tmp_path):
    tpl = tmp_path / "tpl.docx"
    build_synthetic_template(str(tpl))
    row_specs = [
        RowSpec(
            num="1", name="Одиночный", subject="X", inn="1", address="Y",
            floor="1", room="-", date="01.01.2026", object_id="1",
        )
    ]
    out = tmp_path / "out.docx"
    render_letter(str(tpl), _base_context(), row_specs, str(out))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "vMerge" not in xml
