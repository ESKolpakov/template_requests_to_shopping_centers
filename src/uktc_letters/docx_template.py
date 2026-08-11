"""Заполнение шаблона письма (.docx) переменными полями с сохранением
исходной вёрстки, шрифтов и стилей 1-в-1 (ТЗ, разделы 2 и 5.1).

Подход: шаблон НЕ пересобирается заново через API python-docx (создание
параграфов/таблиц с нуля неизбежно теряет мелкие детали вёрстки), а
редактируется на месте — плейсхолдеры вида `{{ИМЯ}}` заменяются прямо
внутри существующих текстовых узлов (`<w:t>`) документа, поэтому все
шрифты, отступы, границы таблиц и т.п. остаются ровно такими, как в
файле, который загрузил пользователь через раздел «Шаблоны».

Требования к самому файлу шаблона (заполняются пользователем в Word,
см. docs/TEMPLATE_GUIDE.md):

  Простые текстовые плейсхолдеры — вставляются как обычный текст в нужном
  месте документа (в адресном блоке, в теле письма, в подписи):
    {{RECIPIENT_BLOCK}}          {{GREETING}}         {{TC_NAME}}
    {{RECIPIENT_BLOCK_MAIN}}     {{TC_ADDRESS}}        {{DEADLINE_DATE}}
    {{RECIPIENT_BLOCK_ADDRESS}}  {{EXECUTOR_NAME}}    {{EXECUTOR_PHONE}}
    {{SIGNATORY_NAME}}           {{SUBJECT_WORD_FORM}} {{VERB_OSUSHESTVLYAL}}
    {{OBJECT_WORD}}              {{TRADE_OBJECT_PHRASE}}
  (RECIPIENT_BLOCK_MAIN/RECIPIENT_BLOCK_ADDRESS — та же информация, что и
  в RECIPIENT_BLOCK, но раздельно на два абзаца с разным отступом —
  см. docs/TEMPLATE_GUIDE.md.)

  Таблица объектов торговли — в шаблоне должна быть готовая таблица с
  заголовком (как в образцах) и ОДНОЙ строкой-образцом данных сразу под
  ним, в ячейках которой стоят плейсхолдеры:
    {{ROW.NUM}} {{ROW.NAME}} {{ROW.SUBJECT}} {{ROW.INN}} {{ROW.ADDRESS}}
    {{ROW.FLOOR}} {{ROW.ROOM}} {{ROW.DATE}} {{ROW.ID}}
  (10-й столбец — «Торговая деятельность подтверждается/не
  подтверждается» — оставляется пустым, без плейсхолдера). Эта
  строка-образец при генерации клонируется по числу строк таблицы письма
  и удаляется из итогового документа.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

__all__ = [
    "RowSpec",
    "ROW_CELL_TOKENS",
    "MERGE_COLUMN_INDEXES",
    "replace_placeholders",
    "fill_objects_table",
    "find_row_template",
    "render_letter",
    "TemplateError",
]


class TemplateError(Exception):
    """Ошибка в структуре загруженного шаблона (не найден нужный
    плейсхолдер/таблица) — показывается пользователю в интерфейсе."""


# ---------------------------------------------------------------------------
# Замена простых плейсхолдеров (возможно, с переносами строк) в тексте.
# ---------------------------------------------------------------------------


def _iter_all_paragraphs(document: Document):
    """Все параграфы документа, включая вложенные в таблицы (любой
    глубины), а также в колонтитулах."""
    for p_elm in document.element.body.iter(qn("w:p")):
        yield Paragraph(p_elm, document)
    for section in document.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                yield p
            for table in part.tables:
                yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _clear_run_content(run) -> None:
    """Удалить у run весь текстовый контент (<w:t>/<w:br>/<w:tab>/<w:cr>),
    сохранив <w:rPr> (форматирование) и сам элемент run."""
    r = run._element
    for child in list(r):
        if child.tag in (qn("w:t"), qn("w:br"), qn("w:tab"), qn("w:cr")):
            r.remove(child)


def _set_run_text_multiline(run, text: str) -> None:
    """Записать в run текст, поддерживая переносы строк (каждая "\\n" в
    `text` становится отдельным <w:br/>) — используется для
    {{RECIPIENT_BLOCK}} (адресный блок в несколько строк)."""
    _clear_run_content(run)
    r = run._element
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            br = r.makeelement(qn("w:br"), {})
            r.append(br)
        t = r.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
        t.text = line
        r.append(t)


def _replace_token_once(paragraph: Paragraph, token: str, value: str) -> bool:
    """Заменить первое вхождение `token` в тексте параграфа на `value`
    (с поддержкой многострочных значений), даже если токен разбит Word'ом
    на несколько runs. Возвращает True, если замена произведена."""
    runs = paragraph.runs
    if not runs:
        return False
    texts = [r.text or "" for r in runs]
    full = "".join(texts)
    idx = full.find(token)
    if idx == -1:
        return False
    end = idx + len(token)

    offsets = []
    pos = 0
    for t in texts:
        offsets.append((pos, pos + len(t)))
        pos += len(t)

    def run_at(char_pos: int) -> int:
        for i, (s, e) in enumerate(offsets):
            if s <= char_pos < e:
                return i
        return len(runs) - 1  # конец текста — последний run

    start_run = run_at(idx)
    end_run = run_at(end - 1)

    start_off = idx - offsets[start_run][0]
    end_off = end - offsets[end_run][0]

    prefix = texts[start_run][:start_off]
    suffix = texts[end_run][end_off:]

    if start_run == end_run:
        _set_run_text_multiline(runs[start_run], prefix + value + suffix)
    else:
        _set_run_text_multiline(runs[start_run], prefix + value)
        for i in range(start_run + 1, end_run):
            _clear_run_content(runs[i])
        _set_run_text_multiline(runs[end_run], suffix)
    return True


def replace_placeholders(document: Document, context: dict[str, str]) -> set[str]:
    """Заменить все плейсхолдеры `{{KEY}}` из `context` во всём документе
    (основной текст, таблицы, колонтитулы). Значения `None` заменяются на
    пустую строку. Возвращает множество ключей, которые были реально
    найдены и заменены хотя бы один раз (пригодится для предупреждений в
    интерфейсе о "неиспользуемых"/отсутствующих плейсхолдерах)."""
    found: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        for key, value in context.items():
            token = "{{" + key + "}}"
            text_value = "" if value is None else str(value)
            # Токен может встретиться несколько раз в одном параграфе —
            # заменяем, пока находим (с защитой от бесконечного цикла).
            for _ in range(50):
                if _replace_token_once(paragraph, token, text_value):
                    found.add(key)
                else:
                    break
    return found


# ---------------------------------------------------------------------------
# Таблица объектов торговли: клонирование строки-образца + rowspan.
# ---------------------------------------------------------------------------

ROW_CELL_TOKENS = [
    "{{ROW.NUM}}",
    "{{ROW.NAME}}",
    "{{ROW.SUBJECT}}",
    "{{ROW.INN}}",
    "{{ROW.ADDRESS}}",
    "{{ROW.FLOOR}}",
    "{{ROW.ROOM}}",
    "{{ROW.DATE}}",
    "{{ROW.ID}}",
]

# Индексы столбцов (считая с 0), которые объединяются по вертикали
# (rowspan), когда несколько строк относятся к одному "Наименованию
# объекта" (ТЗ, раздел 5): Наименование объекта(1), Адрес(4), Этаж(5),
# Помещение(6), Дата сбора информации(7), ID(8). Не объединяются:
# № п/п(0), Хозяйствующий субъект(2), ИНН(3), отметка(9).
MERGE_COLUMN_INDEXES = {1, 4, 5, 6, 7, 8}


@dataclass
class RowSpec:
    """Одна строка таблицы-приложения к письму."""

    num: str
    name: str
    subject: str
    inn: str
    address: str
    floor: str
    room: str
    date: str
    object_id: str
    # True, если эта строка продолжает группу предыдущей строки (те же
    # наименование/адрес/этаж/помещение/дата/ID) — для неё в объединяемых
    # столбцах ставится vMerge-continue, а видимый текст пуст.
    merge_continue: bool = False
    # True, если это первая строка группы из >1 строк — для неё ставится
    # vMerge-restart в объединяемых столбцах.
    merge_start: bool = False

    def token_values(self) -> list[str]:
        return [
            self.num,
            self.name,
            self.subject,
            self.inn,
            self.address,
            self.floor,
            self.room,
            self.date,
            self.object_id,
        ]


def build_row_specs(rows: list) -> list[RowSpec]:
    """Построить список RowSpec из строк группы объектов одного здания
    (`objects_source.TradeObjectRow`), проставив флаги объединения ячеек
    для одинаковых "объектов" (см. `TradeObjectRow.group_fields`)."""
    specs: list[RowSpec] = []
    prev_group_fields = None
    group_start_index = -1
    for i, row in enumerate(rows):
        spec = RowSpec(
            num=str(i + 1),
            name=row.object_name,
            subject=row.subject,
            inn=row.inn,
            address=row.building_address,
            floor=row.floor,
            room=row.room_number,
            date=row.survey_date,
            object_id=row.object_id,
        )
        if row.group_fields == prev_group_fields:
            spec.merge_continue = True
            specs[group_start_index].merge_start = True
        else:
            group_start_index = i
            prev_group_fields = row.group_fields
        specs.append(spec)
    return specs


def find_row_template(document: Document) -> tuple[Table, int]:
    """Найти в документе таблицу и индекс строки-образца (первая ячейка
    которой содержит токен {{ROW.NUM}})."""
    marker = ROW_CELL_TOKENS[0]
    for table in document.tables:
        for row_idx, row in enumerate(table.rows):
            if not row.cells:
                continue
            if marker in row.cells[0].text:
                return table, row_idx
    raise TemplateError(
        f"В шаблоне не найдена строка-образец таблицы объектов торговли "
        f"(ячейка с плейсхолдером {marker})."
    )


def _set_vmerge(cell: _Cell, mode: str | None) -> None:
    """mode: "restart" | "continue" | None (убрать объединение)."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:vMerge"))
    if existing is not None:
        tcPr.remove(existing)
    if mode is None:
        return
    vmerge = tcPr.makeelement(qn("w:vMerge"), {})
    if mode == "restart":
        vmerge.set(qn("w:val"), "restart")
    tcPr.append(vmerge)


def _clear_cell_text(cell: _Cell) -> None:
    for paragraph in cell.paragraphs:
        for run in list(paragraph.runs):
            _clear_run_content(run)


def fill_objects_table(document: Document, row_specs: list[RowSpec]) -> None:
    """Заполнить таблицу объектов торговли: клонировать строку-образец по
    числу `row_specs`, подставить значения и объединить ячейки для строк
    одной группы (rowspan), затем удалить исходную строку-образец."""
    table, template_row_idx = find_row_template(document)
    template_tr = table.rows[template_row_idx]._tr
    tbl = table._tbl

    new_trs = []
    for spec in row_specs:
        tr_copy = copy.deepcopy(template_tr)
        tbl.insert(list(tbl).index(template_tr) + 1 + len(new_trs), tr_copy)
        new_trs.append(tr_copy)

    # После вставки клонов работаем с ними через свежий объект Table,
    # т.к. добавление <w:tr> напрямую в XML не обновляет закэшированный
    # список строк python-docx.
    table = Table(tbl, table._parent)
    start = template_row_idx + 1
    for offset, spec in enumerate(row_specs):
        row = table.rows[start + offset]
        cells = row.cells
        for col_idx, token in enumerate(ROW_CELL_TOKENS):
            if col_idx >= len(cells):
                break
            cell = cells[col_idx]
            for paragraph in cell.paragraphs:
                _replace_token_once(paragraph, token, spec.token_values()[col_idx])

        for col_idx in MERGE_COLUMN_INDEXES:
            if col_idx >= len(cells):
                continue
            cell = cells[col_idx]
            if spec.merge_continue:
                _set_vmerge(cell, "continue")
                _clear_cell_text(cell)
            elif spec.merge_start:
                _set_vmerge(cell, "restart")
            else:
                _set_vmerge(cell, None)

    # Удаляем исходную строку-образец.
    tbl.remove(template_tr)


# ---------------------------------------------------------------------------
# Высокоуровневая функция рендера письма.
# ---------------------------------------------------------------------------


def render_letter(template_path: str, context: dict[str, str], row_specs: list[RowSpec], output_path: str) -> set[str]:
    """Открыть шаблон, заполнить простые плейсхолдеры и таблицу объектов,
    сохранить результат в `output_path`. Возвращает множество реально
    найденных и заменённых ключей `context` (для диагностики в
    интерфейсе)."""
    document = Document(template_path)
    fill_objects_table(document, row_specs)
    found = replace_placeholders(document, context)
    document.save(output_path)
    return found
