"""Разбор "сырой" выгрузки объектов торговли ("Реестр ТО" из АИС ОПН, см.
ТЗ п.3.1) и группировка строк по зданию для формирования одного письма на
адрес.

Файл загружается КАК ЕСТЬ, без предварительной ручной обработки:
- строка 1 — заголовок отчёта ("Реестр ТО (сформировано ... )"),
- строка 2 — реальные заголовки столбцов,
- начиная со строки 3 — данные.

Из 56 столбцов используются только 9 (см. таблицу в ТЗ), остальные
игнорируются полностью.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

from .addresses import building_key, split_room

__all__ = [
    "TradeObjectRow",
    "BuildingGroup",
    "clean_object_name",
    "load_trade_objects",
    "group_by_building",
]

# Заголовки нужных столбцов, как они выглядят в исходном файле (строка 2).
# Сопоставление идёт по названию столбца, а не по номеру/букве — так
# устойчивее к добавлению/перестановке столбцов АИС ОПН в будущих выгрузках.
_COLUMN_MAP = {
    "№ п/п": "seq_no",
    "Наименование ТО": "object_name_raw",
    "Хозяйствующий субъект": "subject",
    "ИНН": "inn",
    "Адрес ТО": "address_raw",
    "Этаж": "floor",
    "№ квартиры (офиса)": "room_number",
    "Дата последнего обхода/АБО": "survey_date",
    "Системный идентификатор объекта": "object_id",
}

_PARENS_RE = re.compile(r"\s*\([^)]*\)")


def clean_object_name(raw_name: str) -> str:
    """Убрать из названия объекта всё содержимое в скобках вместе со
    скобками: "Мастер ПОЛО (ТЦ до 14.08.2026)" -> "Мастер ПОЛО".
    """
    if raw_name is None:
        return ""
    name = _PARENS_RE.sub("", str(raw_name))
    return re.sub(r"\s+", " ", name).strip()


@dataclass
class TradeObjectRow:
    """Одна строка выгрузки после извлечения нужных полей."""

    seq_no: str
    object_name: str
    subject: str
    inn: str
    address_raw: str
    floor: str
    room_number: str
    survey_date: str
    object_id: str
    building_address: str  # адрес без "комната ..."

    @property
    def group_fields(self) -> tuple:
        """Поля, по которым объединяются строки одного "Наименования
        объекта" (rowspan в таблице письма — см. ТЗ п.3.1/5)."""
        return (
            self.object_name,
            self.building_address,
            self.floor,
            self.room_number,
            self.survey_date,
            self.object_id,
        )


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(ws) -> int:
    """Найти строку с реальными заголовками столбцов (обычно строка 2 —
    строка 1 занята заголовком отчёта "Реестр ТО (сформировано ...)")."""
    for row_idx in range(1, min(ws.max_row, 5) + 1):
        values = {
            _cell_str(ws.cell(row=row_idx, column=c).value)
            for c in range(1, ws.max_column + 1)
        }
        if "№ п/п" in values and "Наименование ТО" in values:
            return row_idx
    raise ValueError(
        "Не удалось найти строку заголовков столбцов в выгрузке объектов "
        "торговли (ожидались столбцы «№ п/п», «Наименование ТО», ...)."
    )


def load_trade_objects(path: str, sheet_name: str | None = None) -> list[TradeObjectRow]:
    """Загрузить и разобрать выгрузку "Реестр ТО" из файла `path`.

    `sheet_name` — опциональное имя листа; по умолчанию берётся первый лист
    книги (в образце — единственный содержательный лист "Реестр ТО").
    """
    # read_only=True не используется: в реальных выгрузках АИС ОПН
    # встречаются файлы с некорректно объявленными границами листа
    # (dimension), из-за чего openpyxl в режиме read_only занижает
    # max_row/max_column и "теряет" данные.
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    header_row = _find_header_row(ws)
    col_by_name: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        title = _cell_str(ws.cell(row=header_row, column=c).value)
        if title in _COLUMN_MAP:
            col_by_name[_COLUMN_MAP[title]] = c

    missing = set(_COLUMN_MAP.values()) - set(col_by_name)
    if missing:
        raise ValueError(
            "В выгрузке объектов торговли не найдены обязательные столбцы: "
            + ", ".join(sorted(missing))
        )

    rows: list[TradeObjectRow] = []
    for r in range(header_row + 1, ws.max_row + 1):
        seq_no = _cell_str(ws.cell(row=r, column=col_by_name["seq_no"]).value)
        address_raw = _cell_str(ws.cell(row=r, column=col_by_name["address_raw"]).value)
        object_name_raw = _cell_str(ws.cell(row=r, column=col_by_name["object_name_raw"]).value)
        # Пустая строка (нет ни номера, ни адреса, ни названия) — пропускаем
        # (защита от строк-разделителей / хвостовых пустых строк листа).
        if not seq_no and not address_raw and not object_name_raw:
            continue
        building_address, _room = split_room(address_raw)
        rows.append(
            TradeObjectRow(
                seq_no=seq_no,
                object_name=clean_object_name(object_name_raw),
                subject=_cell_str(ws.cell(row=r, column=col_by_name["subject"]).value),
                inn=_cell_str(ws.cell(row=r, column=col_by_name["inn"]).value),
                address_raw=address_raw,
                floor=_cell_str(ws.cell(row=r, column=col_by_name["floor"]).value),
                room_number=_cell_str(ws.cell(row=r, column=col_by_name["room_number"]).value),
                survey_date=_cell_str(ws.cell(row=r, column=col_by_name["survey_date"]).value),
                object_id=_cell_str(ws.cell(row=r, column=col_by_name["object_id"]).value),
                building_address=building_address,
            )
        )
    return rows


@dataclass
class BuildingGroup:
    """Группа строк выгрузки, относящихся к одному зданию (одному письму)."""

    key: str  # нормализованный адрес здания (building_key)
    building_address: str  # адрес здания "как в источнике" (первое вхождение)
    rows: list[TradeObjectRow] = field(default_factory=list)


def group_by_building(rows: list[TradeObjectRow]) -> list[BuildingGroup]:
    """Сгруппировать строки выгрузки по зданию (адрес до ", комната").

    Порядок групп — по первому появлению адреса в исходном файле; порядок
    строк внутри группы сохраняется как в источнике.
    """
    groups: dict[str, BuildingGroup] = {}
    order: list[str] = []
    for row in rows:
        key = building_key(row.building_address)
        if key not in groups:
            groups[key] = BuildingGroup(key=key, building_address=row.building_address)
            order.append(key)
        groups[key].rows.append(row)
    return [groups[k] for k in order]
