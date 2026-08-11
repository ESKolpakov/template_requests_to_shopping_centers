"""Вкладка «Шаблоны» (ТЗ, раздел 2)."""

from __future__ import annotations

import shutil
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from ..docx_template import ROW_CELL_TOKENS, TemplateError, find_row_template
from ..paths import default_templates_dir
from .state import AppState

_REQUIRED_TOKENS = [
    "{{RECIPIENT_BLOCK}}",
    "{{GREETING}}",
    "{{TC_NAME}}",
    "{{TC_ADDRESS}}",
    "{{DEADLINE_DATE}}",
    "{{EXECUTOR_NAME}}",
    "{{EXECUTOR_PHONE}}",
    "{{SIGNATORY_NAME}}",
]


class TemplatesTab(ttk.Frame):
    def __init__(self, parent, state: AppState):
        super().__init__(parent, padding=12)
        self.state = state

        ttk.Label(
            self,
            text=(
                "Шаблон письма — обычный .docx с плейсхолдерами вида "
                "{{ИМЯ}} в нужных местах текста и готовой строкой-образцом "
                "в таблице объектов торговли. Подробности — в "
                "docs/TEMPLATE_GUIDE.md."
            ),
            wraplength=560,
            justify="left",
        ).pack(anchor="w", pady=(0, 10))

        path_frame = ttk.Frame(self)
        path_frame.pack(fill="x", pady=4)
        ttk.Label(path_frame, text="Текущий шаблон:").pack(side="left")
        self.path_var = tk.StringVar(value=self.state.template_path or "(не выбран)")
        ttk.Label(path_frame, textvariable=self.path_var, foreground="#333").pack(
            side="left", padx=6
        )

        btns = ttk.Frame(self)
        btns.pack(fill="x", pady=8)
        ttk.Button(btns, text="Загрузить шаблон...", command=self._load_template).pack(
            side="left"
        )
        ttk.Button(btns, text="Проверить структуру шаблона", command=self._validate).pack(
            side="left", padx=6
        )

        self.status_text = tk.Text(self, height=14, wrap="word")
        self.status_text.pack(fill="both", expand=True, pady=(10, 0))
        self.status_text.configure(state="disabled")

        if self.state.template_path:
            self._validate()

    def _log(self, text: str) -> None:
        self.status_text.configure(state="normal")
        self.status_text.delete("1.0", "end")
        self.status_text.insert("1.0", text)
        self.status_text.configure(state="disabled")

    def _load_template(self) -> None:
        path = filedialog.askopenfilename(
            title="Выбрать шаблон письма", filetypes=[("Word документы", "*.docx")]
        )
        if not path:
            return
        dest_dir = default_templates_dir()
        dest = dest_dir / Path(path).name
        try:
            shutil.copy(path, dest)
        except OSError as exc:
            messagebox.showerror("Ошибка", f"Не удалось скопировать шаблон: {exc}")
            return
        self.state.template_path = str(dest)
        self.state.settings.template_path = str(dest)
        self.state.save_settings()
        self.path_var.set(self.state.template_path)
        self._validate()

    def _validate(self) -> None:
        if not self.state.template_path:
            self._log("Шаблон не выбран.")
            return
        try:
            import docx

            document = docx.Document(self.state.template_path)
        except Exception as exc:  # noqa: BLE001 — показываем пользователю как есть
            self._log(f"Не удалось открыть файл шаблона:\n{exc}")
            return

        lines = ["Шаблон открыт успешно.", ""]

        full_text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        missing_simple = [t for t in _REQUIRED_TOKENS if t not in full_text]
        if missing_simple:
            lines.append("⚠ Не найдены плейсхолдеры: " + ", ".join(missing_simple))
        else:
            lines.append("✓ Все основные текстовые плейсхолдеры на месте.")

        try:
            find_row_template(document)
            lines.append("✓ Найдена строка-образец таблицы объектов торговли.")
        except TemplateError as exc:
            lines.append(f"⚠ {exc}")

        missing_row_tokens = []
        for token in ROW_CELL_TOKENS:
            if token not in full_text:
                missing_row_tokens.append(token)
        if missing_row_tokens:
            lines.append(
                "⚠ Не найдены плейсхолдеры строки таблицы: "
                + ", ".join(missing_row_tokens)
            )

        self._log("\n".join(lines))
